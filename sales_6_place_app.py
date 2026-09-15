import streamlit as st
import pandas as pd
import re
from io import BytesIO
from openpyxl.utils import get_column_letter

PLACE_NAMES = ["Aman", "Betty", "Dagi", "Mahi", "Elsa", "Ekram"]

EXPORT_COLUMNS = [
    "Date", "Month", "Year", "Start time", "End time", "Duration", "Shop no",
    "Owner Name", "Contact", "Landmark", "Geo", "Channel",
    "SVZ Order", "SV Order", "SVJ Order", "Total Order",
    "SVZ Availability", "SV Availability", "SVJ Availability",
    "SVZ Date", "SV Date", "SVJ Date", "Reason",
    "Lat", "Lon", "Distance"
]


def format_time_24h(time_text):
    """Return time as HH:MM only, with no AM/PM."""
    from datetime import datetime
    try:
        return datetime.strptime(time_text.strip().upper(), "%I:%M %p").strftime("%H:%M")
    except ValueError:
        return ""


def extract_start_time(text):
    m = re.search(r"\[(\d{1,2}/\d{1,2}/\d{4})\s+(\d{1,2}:\d{2}\s*[AP]M)\]", text, re.I)
    return format_time_24h(m.group(2)) if m else ""


def extract_date_parts(text):
    m = re.search(r"\[(\d{1,2})/(\d{1,2})/(\d{4})\s+(\d{1,2}:\d{2}\s*[AP]M)\]", text, re.I)
    if not m:
        return "", "", "", ""
    month, day, year, tm = int(m.group(1)), int(m.group(2)), int(m.group(3)), m.group(4)
    # Match the existing Excel layout: Date is the day number, followed by
    # Month and Year as separate columns.
    return day, month, year, format_time_24h(tm)


def extract_field(text, labels):
    """Extract a labeled field and clean WhatsApp bullets/dots."""
    label_pat = "|".join(re.escape(x) for x in labels)
    # Stop at the next known label, including labels with decorative dots.
    stop = r"(?=\s*(?:●|\*|-|•)?\s*(?:Shop|Shope|LShope|Land\s*Mark|Landmark|Owner(?:\s+Name)?|Phone(?:\s+number)?|Phonenumber|customer\s*type|Customer\s*Type|Order|Aval[iy]ablity|Availability|Reason)\b)"
    m = re.search(
        rf"(?:^|\s)(?:●|\*|-|•)?\s*(?:{label_pat})\s*(?:[:.\-]+\s*)?(.*?){stop}|"
        rf"(?:^|\s)(?:●|\*|-|•)?\s*(?:{label_pat})\s*(?:[:.\-]+\s*)?(.*)$",
        text, re.I | re.S
    )
    value = (m.group(1) if m and m.group(1) is not None else (m.group(2) if m else ""))
    # Remove decorative WhatsApp bullets and leading/trailing dots/spaces.
    value = re.sub(r"^[\s●•*\-.:]+", "", value)
    value = re.sub(r"[\s●•*\-.:]+$", "", value)
    value = re.sub(r"\s+", " ", value)
    return value.strip()


def extract_phone(text):
    m = re.search(r"(?:Phone\s*number|Phonenumber|Phone)\s*[:.\-]?\s*([0-9][0-9\s\-]{6,})", text, re.I)
    return re.sub(r"\D", "", m.group(1)) if m else ""


def extract_shop_no(text):
    m = re.search(r"(?:^|\s)(?:L\s*)?(?:Shop|Shope)\s*No\s*[:.\-]?\s*([A-Za-z0-9]+)", text, re.I)
    return m.group(1).strip() if m else ""



def extract_geo(text):
    """Extract lat/lon from a pasted Google Maps URL or plain coordinates."""
    if not text:
        return "", "", ""

    candidates = [
        # q=9.020863,38.840807
        r"(?:[?&](?:q|ll)=)(-?\d+(?:\.\d+)?)[,%]2?C?\s*(-?\d+(?:\.\d+)?)",
        # Explicit latitude/longitude URL parameters
        r"(?:latitude|lat)=(-?\d+(?:\.\d+)?)[^0-9\-]+(?:longitude|lon|lng)=(-?\d+(?:\.\d+)?)",
        # Coordinates anywhere in pasted URL/text
        r"(?<!\d)(-?\d{1,3}\.\d{4,})\s*[,;]\s*(-?\d{1,3}\.\d{4,})(?!\d)",
    ]

    for pat in candidates:
        m = re.search(pat, text, re.I)
        if m:
            lat = float(m.group(1))
            lon = float(m.group(2))
            if -90 <= lat <= 90 and -180 <= lon <= 180:
                return f"{lat},{lon}", lat, lon

    return "", "", ""


def haversine_km(lat1, lon1, lat2, lon2):
    import math
    if lat1 == "" or lon1 == "" or lat2 == "" or lon2 == "":
        return ""
    R = 6371.0
    p1, p2 = math.radians(float(lat1)), math.radians(float(lat2))
    dp = math.radians(float(lat2) - float(lat1))
    dl = math.radians(float(lon2) - float(lon1))
    a = math.sin(dp/2)**2 + math.cos(p1) * math.cos(p2) * math.sin(dl/2)**2
    return round(2 * R * math.asin(math.sqrt(a)), 2)



def normalize_number(value):
    if not value:
        return ""
    value = value.strip()
    try:
        return "" if float(value) == 0 else value
    except ValueError:
        return value


def find_section(text, start_words, end_words):
    """Extract a section while tolerating misspelled labels and labels on separate lines."""
    starts = "|".join(re.escape(x) for x in start_words)
    ends = "|".join(re.escape(x) for x in end_words)
    m = re.search(
        rf"(?:^|\s)(?:{starts})\s*[:.\-]*\s*(.*?)(?=(?:^|\s)(?:{ends})\s*[:.\-]*|$)",
        text, re.I | re.S
    )
    return m.group(1) if m else ""


def extract_product(section, product):
    """
    Robustly read quantities from strings such as:
      SVz40pc
      LSV0ctn
      SVJ80pc
      SVZ.....0 pcs
      SV.......0 ctn
      Svg0cartons
      Sz:0carton

    For SV, LSV/Sz/Svg are accepted aliases, but SVZ/SVJ are excluded.
    """
    if not section:
        return ""

    if product == "SVZ":
        aliases = r"SVZ"
    elif product == "SVJ":
        aliases = r"SVJ"
    else:
        aliases = r"(?:LSV|SVG|SZ|SV)(?!Z|J)"

    # IMPORTANT: do not require a word boundary after the label because
    # real reports use forms like SVz40pc, LSV0ctn, SVJ80pc, SVG0cartons.
    m = re.search(
        rf"(?<![A-Za-z]){aliases}\s*[:.\-]*\s*(\d+(?:\.\d+)?)",
        section, re.I
    )
    return normalize_number(m.group(1)) if m else ""


def parse_availability_lines(text):
    """
    Parse availability directly from the full message. This is deliberately
    independent of the spelling/formatting of the 'Availability' heading.
    """
    # Only inspect text after the first Availability/Avaliablity-like heading.
    m = re.search(r"\bAva[iy]labl?i?t?y\b|\bAvailability\b|\bAvaliablity\b", text, re.I)
    if not m:
        return "", "", ""

    tail = text[m.end():]
    # Stop before Reason.
    tail = re.split(r"\bReason\b", tail, maxsplit=1, flags=re.I)[0]

    return (
        extract_product(tail, "SVZ"),
        extract_product(tail, "SV"),
        extract_product(tail, "SVJ"),
    )


def parse_order_lines(text):
    m = re.search(r"\bOrder\b", text, re.I)
    if not m:
        return "", "", ""
    tail = text[m.end():]
    tail = re.split(
        r"\bAva[iy]labl?i?t?y\b|\bAvailability\b|\bAvaliablity\b|\bReason\b",
        tail, maxsplit=1, flags=re.I
    )[0]
    return (
        extract_product(tail, "SVZ"),
        extract_product(tail, "SV"),
        extract_product(tail, "SVJ"),
    )



def parse_shop(shop_text, location_text, base_lat, base_lon, place):
    date, month, year, start_time = extract_date_parts(shop_text)

    # Keep the FULL pasted Google Maps link in Geo, while separately
    # extracting latitude/longitude for calculations.
    geo_source = location_text.strip() if location_text and location_text.strip() else shop_text
    geo, lat, lon = extract_geo(geo_source)
    if location_text and location_text.strip():
        geo_value = location_text.strip()
        # If the pasted location is only coordinates, make a usable Maps URL.
        if not re.search(r"https?://", geo_value, re.I) and lat != "" and lon != "":
            geo_value = f"https://maps.google.com/maps?q={lat},{lon}&ll={lat},{lon}&z=16"
    else:
        geo_value = geo

    customer = extract_field(shop_text, ["customer type", "Customer Type"]).lower()
    channel = (
        "retail" if re.match(r"r(?:etail)?\b", customer)
        else ("wholesale" if re.match(r"w(?:holesale)?\b", customer) else "")
    )

    svz_order, sv_order, svj_order = parse_order_lines(shop_text)
    svz_av, sv_av, svj_av = parse_availability_lines(shop_text)

    reason_m = re.search(
        r"\bReason\b\s*[:.\-]*\s*(.*?)(?=\[?\s*Location\s*\]?|$)",
        shop_text, re.I | re.S
    )
    reason = reason_m.group(1) if reason_m else ""
    reason = re.sub(r"^[\s●•*\-.:]+|[\s●•*\-.:]+$", "", reason)
    reason = re.sub(r"\s+", " ", reason).strip()

    return {
        "Date": date,
        "Month": month,
        "Year": year,
        "Start time": start_time,
        "End time": "",
        "Duration": "",
        "Shop no": int(extract_shop_no(shop_text)) if extract_shop_no(shop_text).isdigit() else extract_shop_no(shop_text),
        "Owner Name": extract_field(shop_text, ["Owner Name", "Owner"]),
        "Contact": extract_phone(shop_text),
        "Landmark": extract_field(shop_text, ["Land Mark", "Landmark"]),
        "Geo": geo_value,
        "Channel": channel,
        "SVZ Order": svz_order,
        "SV Order": sv_order,
        "SVJ Order": svj_order,
        "SVZ Availability": svz_av,
        "SV Availability": sv_av,
        "SVJ Availability": svj_av,
        "SVZ Date": "",
        "SV Date": "",
        "SVJ Date": "",
        "Reason": reason,
        "Lat": lat,
        "Lon": lon,
        "Distance": haversine_km(base_lat, base_lon, lat, lon),
        "Place": place,
    }


def split_shops(text):
    """Split WhatsApp export into timestamped message blocks."""
    matches = list(re.finditer(
        r"(?m)(?=\[\d{1,2}/\d{1,2}/\d{4}\s+\d{1,2}:\d{2}\s*[AP]M\])",
        text, re.I
    ))
    if not matches:
        return [(text, "")]

    blocks = []
    for i, m in enumerate(matches):
        end = matches[i + 1].start() if i + 1 < len(matches) else len(text)
        blocks.append((text[m.start():end], ""))
    return blocks


def process_place(text, place, base_lat, base_lon):
    rows = []
    blocks = split_shops(text)

    for i, (block, _) in enumerate(blocks):
        if not re.search(r"(?:Shop|Shope|LShope)\s*No", block, re.I):
            continue

        # First look for [ Location ] inside this block.
        loc_match = re.search(
            r"\[?\s*Location\s*\]?\s*(.*?)(?=\[\d{1,2}/\d{1,2}/\d{4}\s+\d{1,2}:\d{2}\s*[AP]M\]|$)",
            block, re.I | re.S
        )
        location = loc_match.group(1) if loc_match else ""
        shop_part = block[:loc_match.start()] if loc_match else block

        # If Location is the next WhatsApp message, it may be in the next block.
        if not location and i + 1 < len(blocks):
            next_block = blocks[i + 1][0]
            if re.search(r"\[?\s*Location\s*\]?", next_block, re.I):
                lm = re.search(r"\[?\s*Location\s*\]?\s*(.*)", next_block, re.I | re.S)
                location = lm.group(1) if lm else ""

        rows.append(parse_shop(shop_part, location, base_lat, base_lon, place))

    return rows


def add_end_times_and_durations(df):
    """Use the next shop's start time as the current shop's end time.
    Duration is calculated only when the next shop is on the same date.
    """
    if df is None or df.empty:
        return df
    df = df.copy()
    df["End time"] = ""
    df["Duration"] = ""

    from datetime import datetime
    for place in df["Place"].dropna().unique():
        idxs = list(df.index[df["Place"] == place])
        for pos, idx in enumerate(idxs[:-1]):
            nxt = idxs[pos + 1]
            if str(df.at[idx, "Date"]) != str(df.at[nxt, "Date"]):
                continue
            st = str(df.at[idx, "Start time"]).strip()
            et = str(df.at[nxt, "Start time"]).strip()
            if not st or not et:
                continue
            try:
                t1 = datetime.strptime(st, "%H:%M")
                t2 = datetime.strptime(et, "%H:%M")
                delta = t2 - t1
                if delta.total_seconds() < 0:
                    delta = delta.replace(days=1)
                mins = int(delta.total_seconds() // 60)
                df.at[idx, "End time"] = et
                df.at[idx, "Duration"] = f"{mins // 60}:{mins % 60:02d}"
            except ValueError:
                pass
    return df


def add_total_order(df):
    """Add Total Order immediately after SVJ Order.
    Total = SVZ Order + SV Order + SVJ Order; blanks are treated as zero.
    """
    if df is None or df.empty:
        return df
    df = df.copy()

    def num(v):
        try:
            if v is None or str(v).strip() == "":
                return 0.0
            return float(str(v).replace(",", "").strip())
        except (ValueError, TypeError):
            return 0.0

    df["Total Order"] = [
        num(a) + num(b) + num(c)
        for a, b, c in zip(df["SVZ Order"], df["SV Order"], df["SVJ Order"])
    ]
    # Show whole numbers as integers when possible.
    df["Total Order"] = df["Total Order"].apply(
        lambda x: "" if float(x) == 0 else (int(x) if float(x).is_integer() else x)
    )
    return df


def process_all(place_texts, base_lat, base_lon):
    rows = []
    for place in PLACE_NAMES:
        text = place_texts.get(place, "")
        if text and text.strip():
            rows.extend(process_place(text, place, base_lat, base_lon))
    if not rows:
        return pd.DataFrame(columns=EXPORT_COLUMNS + ["Place"])
    return add_total_order(add_end_times_and_durations(pd.DataFrame(rows)))


def append_to_history(history_df, new_df):
    if new_df is None or new_df.empty:
        return history_df.copy()
    if history_df is None or history_df.empty:
        combined = new_df.copy()
    else:
        combined = pd.concat([history_df, new_df], ignore_index=True)

    combined = combined.drop_duplicates(
        subset=["Date", "Start time", "Shop no", "Owner Name", "Contact", "Geo", "Place"],
        keep="first"
    ).reset_index(drop=True)
    # Recalculate end times/durations across the complete accumulated history,
    # so adding a new shop immediately fills the previous shop's end time.
    return add_total_order(add_end_times_and_durations(combined))


def make_workbook(all_df):
    all_df = add_total_order(all_df)
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        for person in PLACE_NAMES:
            df = all_df[all_df["Place"].astype(str).str.lower() == person.lower()].copy()
            if "Place" in df.columns:
                df = df.drop(columns=["Place"])
            for col in EXPORT_COLUMNS:
                if col not in df.columns:
                    df[col] = ""
            df = df[EXPORT_COLUMNS]
            df.to_excel(writer, index=False, sheet_name=person)

            # Format the workbook to look like a clean, practical field-sales
            # sheet similar to the user's existing Excel screenshot.
            ws = writer.book[person]
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            ws.sheet_view.showGridLines = True

            from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
            header_fill = PatternFill(fill_type="solid", fgColor="D9E1F2")
            header_font = Font(bold=True)
            thin = Side(style="thin", color="A6A6A6")
            border = Border(left=thin, right=thin, top=thin, bottom=thin)

            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.border = border
                cell.alignment = Alignment(horizontal="center", vertical="center")
            ws.row_dimensions[1].height = 22

            # Practical widths: narrow numeric fields, wider names/landmarks,
            # and a wide Geo column so the full pasted Maps URL remains visible.
            widths = {
                "A": 10, "B": 9, "C": 9, "D": 12, "E": 12, "F": 11,
                "G": 10, "H": 20, "I": 16, "J": 18, "K": 55, "L": 13,
                "M": 13, "N": 13, "O": 13, "P": 15, "Q": 18, "R": 18,
                "S": 18, "T": 12, "U": 12, "V": 12, "W": 30, "X": 12,
                "Y": 12, "Z": 12
            }
            for col_letter, width in widths.items():
                ws.column_dimensions[col_letter].width = width

            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.border = border
                    cell.alignment = Alignment(vertical="center")
                # Keep long URLs readable and clickable.
                geo_cell = row[10]
                if isinstance(geo_cell.value, str) and geo_cell.value.startswith("http"):
                    geo_cell.hyperlink = geo_cell.value
                    geo_cell.style = "Hyperlink"

            # Make numeric columns numeric where possible.
            for col in ["B", "C", "G", "M", "N", "O", "P", "Q", "R", "S", "T", "U", "V", "X", "Y", "Z"]:
                for cell in ws[col][1:]:
                    if isinstance(cell.value, str):
                        try:
                            cell.value = float(cell.value) if "." in cell.value else int(cell.value)
                        except (ValueError, TypeError):
                            pass

    output.seek(0)
    return output


st.set_page_config(page_title="WhatsApp Sales → Excel", layout="wide")
st.title("WhatsApp Sales → Excel")
st.caption("Paste reports for all six people. New rows append immediately. The Excel export follows your existing sheet layout: Date = day, Month, Year, Start/End time (24-hour HH:MM, no AM/PM), Duration, then shop/order/availability details. Zero order/availability values are shown as blank, and Geo keeps the full Google Maps link.")

if "history_df" not in st.session_state:
    st.session_state.history_df = pd.DataFrame()

with st.sidebar:
    st.header("Distance base")
    base_lat = st.number_input("Base latitude", value=9.03, format="%.6f")
    base_lon = st.number_input("Base longitude", value=38.74, format="%.6f")
    st.divider()
    st.write("Excel sheets:")
    for person in PLACE_NAMES:
        st.write("• " + person)

st.subheader("Paste the six reports")
place_texts = {}
for i, person in enumerate(PLACE_NAMES, 1):
    place_texts[person] = st.text_area(
        f"{i}. {person}",
        height=220,
        key=f"input_{person}",
        placeholder="Paste one or many WhatsApp messages here...",
    )

c1, c2, c3 = st.columns(3)
with c1:
    add_rows = st.button("Parse & Add Rows", type="primary", use_container_width=True)
with c2:
    clear_inputs = st.button("Clear Input Boxes", use_container_width=True)
with c3:
    clear_all = st.button("Clear All Recorded Rows", use_container_width=True)

if clear_inputs:
    for person in PLACE_NAMES:
        st.session_state[f"input_{person}"] = ""
    st.rerun()

if clear_all:
    st.session_state.history_df = pd.DataFrame()
    st.rerun()

if add_rows:
    parsed = process_all(place_texts, base_lat, base_lon)
    if parsed.empty:
        st.warning("No shop rows were detected. Check the WhatsApp text.")
    else:
        before = len(st.session_state.history_df)
        st.session_state.history_df = append_to_history(st.session_state.history_df, parsed)
        added = len(st.session_state.history_df) - before
        st.success(f"{added} new row(s) added. Total recorded rows: {len(st.session_state.history_df)}")

history = st.session_state.history_df

st.subheader("Live recorded data")
if history.empty:
    st.info("Paste your reports and click Parse & Add Rows. The rows will appear here immediately.")
else:
    display = history.copy()
    st.dataframe(display, use_container_width=True, height=550)

    counts = history.groupby("Place").size().reindex(PLACE_NAMES, fill_value=0).rename("Rows").to_frame()
    st.write("Rows recorded per person")
    st.dataframe(counts, use_container_width=True)

    st.subheader("Export")
    excel = make_workbook(history)
    st.download_button(
        "Download ONE Excel — 6 People / 6 Sheets",
        data=excel.getvalue(),
        file_name="WhatsApp_Sales_All_6_People.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        type="primary",
        use_container_width=True,
    )
